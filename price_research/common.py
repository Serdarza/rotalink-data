"""Ortak yardımcılar: nazik HTTP indirme, metin çıkarma (HTML/PDF/Word/Excel/görsel OCR),
fiyat rakamı ve tarih ayrıştırma."""

from __future__ import annotations

import hashlib
import io
import re
import threading
import time
from dataclasses import dataclass, field
from datetime import date
from urllib.parse import urljoin, urlparse

import requests
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

UA = "RotaLinkPriceBot/1.0 (+https://github.com/Serdarza/rotalink-data; kamu tesisi resmi fiyat kontrolu)"

RESMI_UZANTILAR = (".gov.tr", ".edu.tr", ".k12.tr", ".bel.tr", ".pol.tr", ".tsk.tr", ".mil.tr",
                   ".gov.ct.tr", ".edu.ct.tr")

FILE_EXT_RE = re.compile(r"\.(pdf|docx?|xlsx?|jpe?g|png|webp)(\?|#|$)", re.I)
PRICE_WORD_RE = re.compile(r"fiyat|ücret|ucret|tarife|konaklama|oda[\s_-]|pansiyon|misafirhane|konukevi", re.I)


RESMI_EK: set[str] = set()  # kaynak kaydındaki onaylı .gov.tr dışı resmî alan adları


def resmi_alan_adi(url: str, ek_izinli: set[str] | None = None) -> bool:
    host = (urlparse(url).hostname or "").lower()
    if host.endswith(RESMI_UZANTILAR):
        return True
    ek = RESMI_EK if ek_izinli is None else ek_izinli
    return any(host == h or host.endswith("." + h) for h in ek)


def year_price_phrase(text: str, year: int) -> bool:
    """Metinde '2026 yılı konaklama ücretleri', 'fiyat listesi (01.01.2026)' gibi yıl+fiyat ifadesi var mı."""
    y = str(year)
    return bool(re.search(rf"(?<!\d){y}(?!\d)[^\n|]{{0,40}}(fiyat|[uü]cret|tarife)|(fiyat|[uü]cret|tarife)[^\n|]{{0,40}}(?<!\d){y}(?!\d)",
                          text or "", re.I))


# --------------------------------------------------------------------------- HTTP

@dataclass
class FetchResult:
    url: str
    ok: bool
    status: int | None = None
    content: bytes = b""
    ctype: str = ""
    etag: str | None = None
    last_modified: str | None = None
    not_modified: bool = False
    error: str | None = None
    final_url: str | None = None

    @property
    def sha256(self) -> str | None:
        return hashlib.sha256(self.content).hexdigest() if self.content else None


@dataclass
class Fetcher:
    """Host başına hız sınırı, zaman aşımı, yeniden deneme + üstel geri çekilme, çalışma içi önbellek."""

    per_host_delay: float = 1.5
    timeout: float = 40.0
    retries: int = 3
    max_bytes: int = 25_000_000
    _cache: dict = field(default_factory=dict)
    _host_last: dict = field(default_factory=dict)
    _host_locks: dict = field(default_factory=dict)
    _lock: threading.Lock = field(default_factory=threading.Lock)

    def __post_init__(self) -> None:
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": UA, "Accept-Language": "tr-TR,tr;q=0.9"})

    def _wait_host(self, host: str) -> threading.Lock:
        with self._lock:
            lk = self._host_locks.setdefault(host, threading.Lock())
        lk.acquire()
        gap = time.monotonic() - self._host_last.get(host, 0)
        if gap < self.per_host_delay:
            time.sleep(self.per_host_delay - gap)
        return lk

    def get(self, url: str, etag: str | None = None, last_modified: str | None = None) -> FetchResult:
        with self._lock:
            if url in self._cache:
                return self._cache[url]
        host = urlparse(url).hostname or ""
        headers = {}
        if etag:
            headers["If-None-Match"] = etag
        if last_modified:
            headers["If-Modified-Since"] = last_modified
        res = FetchResult(url=url, ok=False)
        for attempt in range(self.retries):
            lk = self._wait_host(host)
            try:
                verify = True
                try:
                    r = self.session.get(url, headers=headers, timeout=self.timeout, stream=True, verify=verify)
                except requests.exceptions.SSLError:
                    r = self.session.get(url, headers=headers, timeout=self.timeout, stream=True, verify=False)
                content = r.raw.read(self.max_bytes + 1, decode_content=True) if r.status_code == 200 else b""
                res = FetchResult(
                    url=url, ok=r.status_code in (200, 304), status=r.status_code, content=content[: self.max_bytes],
                    ctype=r.headers.get("content-type", "").lower(), etag=r.headers.get("ETag"),
                    last_modified=r.headers.get("Last-Modified"), not_modified=r.status_code == 304,
                    final_url=r.url,
                )
                if r.status_code in (429, 500, 502, 503, 504):
                    res.ok = False
                    res.error = f"HTTP {r.status_code}"
                else:
                    break
            except Exception as e:  # noqa: BLE001
                res = FetchResult(url=url, ok=False, error=f"{type(e).__name__}: {e}"[:300])
            finally:
                self._host_last[host] = time.monotonic()
                lk.release()
            time.sleep(min(60, 2 ** (attempt + 1)))
        if res.ok and res.status == 200 and not res.content:
            res.ok, res.error = False, "boş yanıt"
        with self._lock:
            self._cache[url] = res
        return res


# ----------------------------------------------------------------- metin çıkarma

_ocr = None
_ocr_lock = threading.Lock()


def ocr_image_bytes(data: bytes) -> str:
    """RapidOCR ile görsel okuma; satırlar konuma göre birleştirilir, hücreler ' | ' ile ayrılır."""
    global _ocr
    with _ocr_lock:
        if _ocr is None:
            from rapidocr_onnxruntime import RapidOCR
            _ocr = RapidOCR()
        res, _ = _ocr(data)
    if not res:
        return ""
    items = []
    for box, text, _score in res:
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        items.append(((min(ys) + max(ys)) / 2, min(xs), max(ys) - min(ys), text))
    items.sort()
    lines, cur, cur_y, cur_h = [], [], None, 0
    for y, x, h, t in items:
        if cur_y is None or abs(y - cur_y) <= max(8, 0.5 * max(h, cur_h)):
            cur.append((x, t))
            cur_y = y if cur_y is None else (cur_y + y) / 2
            cur_h = max(cur_h, h)
        else:
            lines.append(" | ".join(t for _, t in sorted(cur)))
            cur, cur_y, cur_h = [(x, t)], y, h
    if cur:
        lines.append(" | ".join(t for _, t in sorted(cur)))
    return "\n".join(lines)


def detect_kind(url: str, ctype: str) -> str:
    if "pdf" in ctype:
        return "pdf"
    if "image/" in ctype:
        return "gorsel"
    if "spreadsheet" in ctype or "excel" in ctype:
        return "xlsx" if "openxml" in ctype else "xls"
    if "wordprocessing" in ctype:
        return "docx"
    if "msword" in ctype:
        return "doc"
    m = FILE_EXT_RE.search(url)
    if m and "html" not in ctype:
        e = m.group(1).lower()
        return {"jpg": "gorsel", "jpeg": "gorsel", "png": "gorsel", "webp": "gorsel"}.get(e, e)
    return "html"


@dataclass
class Extracted:
    text: str
    kind: str
    links: list[tuple[str, str]] = field(default_factory=list)  # (mutlak url, bağlantı metni)
    ocr_used: bool = False
    error: str | None = None


def extract(url: str, content: bytes, ctype: str) -> Extracted:
    kind = detect_kind(url, ctype)
    try:
        if kind == "html":
            return _extract_html(url, content)
        if kind == "pdf":
            return _extract_pdf(content)
        if kind == "gorsel":
            return Extracted("[OCR]\n" + ocr_image_bytes(content), "gorsel", ocr_used=True)
        if kind == "docx":
            import docx
            d = docx.Document(io.BytesIO(content))
            out = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                out.append("--- TABLO")
                for row in t.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
            return Extracted("\n".join(out), "word")
        if kind == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"=== SHEET {ws.title} (gizli={ws.sheet_state != 'visible'})")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        out.append(" | ".join("" if v is None else str(v) for v in row))
            return Extracted("\n".join(out), "excel")
        if kind == "xls":
            import xlrd
            wb = xlrd.open_workbook(file_contents=content)
            out = []
            for sh in wb.sheets():
                out.append(f"=== SHEET {sh.name}")
                for r in range(sh.nrows):
                    out.append(" | ".join(str(v) for v in sh.row_values(r)))
            return Extracted("\n".join(out), "excel")
        if kind == "doc":
            return Extracted("", "word", error="eski .doc biçimi okunamadı")
    except Exception as e:  # noqa: BLE001
        return Extracted("", kind, error=f"{type(e).__name__}: {e}"[:300])
    return Extracted("", kind, error="desteklenmeyen içerik")


def _extract_html(url: str, content: bytes) -> Extracted:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content, "lxml")
    for el in soup(["script", "style", "noscript"]):
        el.decompose()
    out = []
    t = soup.find("title")
    if t:
        out.append("BAŞLIK: " + t.get_text(" ", strip=True))
    out.append(re.sub(r"\s+", " ", soup.get_text(" ", strip=True)))
    for tb in soup.find_all("table"):
        out.append("--- TABLO")
        for tr in tb.find_all("tr"):
            out.append(" | ".join(re.sub(r"\s+", " ", c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"])))
    links = []
    for a in soup.find_all("a", href=True):
        links.append((urljoin(url, a["href"].strip()), a.get_text(" ", strip=True)[:120]))
    for im in soup.find_all("img", src=True):
        links.append((urljoin(url, im["src"].strip()), (im.get("alt") or "")[:120]))
    return Extracted("\n".join(out), "web", links=links)


def _extract_pdf(content: bytes) -> Extracted:
    import pymupdf
    doc = pymupdf.open(stream=content, filetype="pdf")
    out, ocr = [], False
    for i, page in enumerate(doc):
        txt = page.get_text("text")
        if len(txt.strip()) < 40:
            txt = "[OCR]\n" + ocr_image_bytes(page.get_pixmap(dpi=200).tobytes("png"))
            ocr = True
        out.append(f"=== SAYFA {i + 1}\n{txt}")
        try:
            for t in page.find_tables().tables:
                out.append("--- TABLO")
                for row in t.extract():
                    out.append(" | ".join((c or "").replace("\n", " ") for c in row))
        except Exception:  # noqa: BLE001
            pass
    return Extracted("\n".join(out), "pdf", ocr_used=ocr)


# ------------------------------------------------------------ rakam / tarih

def source_numbers(text: str) -> set[float]:
    """Kaynak metnindeki bütün sayıları farklı Türkçe/İngilizce yazımlarıyla çıkarır
    (1.500,00 / 1,500 / 1 500 / 1500.50 / 1.100.00)."""
    nums: set[float] = set()
    t = text.replace("\xa0", " ")
    for m in re.finditer(r"\d+(?:[.,\s]\d+)*", t):
        raw = m.group(0)
        pieces = {raw} | set(re.split(r"\s+", raw))
        for m2 in re.finditer(r"\d{1,3}(?:\s\d{3})+", raw):
            pieces.add(m2.group(0).replace(" ", ""))
        for part in pieces:
            part = part.strip(" .,")
            if not part:
                continue
            cands = {part}
            if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", part):
                cands.add(part.replace(".", "").replace(",", "."))
            if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", part):
                cands.add(part.replace(",", ""))
            if re.fullmatch(r"\d+,\d{1,2}", part):
                cands.add(part.replace(",", "."))
            if re.fullmatch(r"\d{1,3}(\.\d{3})+\.\d{2}", part):
                cands.add(part[:-3].replace(".", "") + part[-3:])
            for c in cands:
                try:
                    nums.add(float(c))
                except ValueError:
                    pass
    return nums


def tarife_prices(tarife: dict) -> list[float]:
    tables = tarife.get("tablolar") or ([tarife] if tarife.get("satirlar") else [])
    out = []
    for t in tables:
        for r in t.get("satirlar") or []:
            for v in (r.get("fiyatlar") or {}).values():
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    out.append(float(v))
    return out


_DATE_RE = re.compile(r"(\d{1,2})[./](\d{1,2})[./](20\d{2})")


def dates_in(text: str) -> list[date]:
    out = []
    for d, m, y in _DATE_RE.findall(text or ""):
        try:
            out.append(date(int(y), int(m), int(d)))
        except ValueError:
            pass
    return out


def period_end(donem: str | None) -> date | None:
    """'01.07.2026 – 31.12.2026' gibi bir dönemin bitiş tarihi (iki tarih varsa ikincisi)."""
    ds = dates_in(donem or "")
    return ds[1] if len(ds) >= 2 else None


def tarife_expired(tarife: dict, today: date) -> bool:
    """Tarifedeki bütün dönemlerin bitiş tarihi geçmişse True. Bitişi belirtilmemiş dönem varsa False."""
    ends = []
    tables = tarife.get("tablolar") or []
    donemler = [tarife.get("donem")] + [t.get("donem") for t in tables]
    donemler = [d for d in donemler if d]
    if not donemler:
        return False
    for d in donemler:
        e = period_end(d)
        if e is None:
            return False
        ends.append(e)
    return max(ends) < today


def mentions_year(text: str, year: int) -> bool:
    return bool(re.search(rf"(?<!\d){year}(?!\d)", text or ""))


def norm(s: str) -> str:
    tr = str.maketrans("çğıöşüÇĞİÖŞÜâîû", "cgiosuCGIOSUaiu")
    s = (s or "").translate(tr).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def digits(s: str) -> str:
    return re.sub(r"\D", "", s or "")[-7:]
